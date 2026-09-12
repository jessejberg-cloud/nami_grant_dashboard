"""Generate editable Word and readable HTML from the maintained Markdown guides.
Run with the document runtime Python. Render DOCX with render_docx.py --emit_pdf,
review every page, then copy the PDFs to public/ before committing a release.
"""
from pathlib import Path
import re, html
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
def slug(s): return re.sub(r'[^a-z0-9 -]','',s.lower()).replace(' ','-')
def inline(s): return re.sub(r'\*\*(.*?)\*\*',r'<strong>\1</strong>',re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'<a href="\2">\1</a>',html.escape(s)))
def generate(source,stem,web):
 d=Document(); sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
 for n in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
  st=d.styles[n];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 
 for el in list(d.styles.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
 d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(7)
 d.styles['Normal'].paragraph_format.line_spacing=1.08
 d.styles['Title'].font.size=Pt(25);d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(13)
 if source=='QUICKSTART.md':
  d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6)
 body=[];bid=1
 for line in (ROOT/source).read_text().splitlines():
  if not line:continue
  if line.startswith('#'):
   marks,title=line.split(' ',1);level=len(marks);p=d.add_paragraph(title,'Title' if level==1 else f'Heading {level-1}')
   b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(bid));b.set(qn('w:name'),slug(title));p._p.insert(0,b);end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bid));p._p.append(end);bid+=1
   body.append(f'<h{level} id="{slug(title)}">{html.escape(title)}</h{level}>');continue
  m=re.fullmatch(r'\d+\. \[([^]]+)\]\(#([^)]+)\)',line)
  if m:
   p=d.add_paragraph();h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),m[2]);r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=m[1];r.append(t);h.append(r);p._p.append(h);p.paragraph_format.space_after=Pt(2);body.append(f'<p class="toc"><a href="#{m[2]}">{html.escape(m[1])}</a></p>');continue
  style='Normal';text=line
  if line.startswith('- '):style='List Bullet';text=line[2:]
  elif re.match(r'^\d+\. ',line):style='Normal';text=line
  p=d.add_paragraph(style=style)
  for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
   run=p.add_run(part);run.bold=i%2==1
  body.append('<p>'+inline(line)+'</p>')
 footer=sec.footer.paragraphs[0];footer.alignment=2;r=footer.add_run('Nami Grant Workspace  |  ');r.font.size=Pt(9)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 d.core_properties.title='Nami Grant Workspace User Manual' if source=='USER_MANUAL.md' else 'Nami Grant Workspace Quick Start'
 d.core_properties.author='Nami Grant Workspace';d.save(ROOT/'public'/f'{stem}.docx')
 (ROOT/'public'/web).write_text('<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+d.core_properties.title+'</title><style>body{font:17px/1.65 Arial,sans-serif;color:#192f39;max-width:900px;margin:auto;padding:24px}a{color:#126b76}h1,h2,h3{line-height:1.3}h2{margin-top:2em}p{overflow-wrap:anywhere}.toc{margin:4px 0}nav{display:flex;flex-wrap:wrap;gap:20px}a:focus{outline:3px solid #126b76}</style><nav><a href="/">Return to dashboard</a><a href="/'+stem+'.docx">Word download</a><a href="/'+stem+'.pdf">PDF download</a></nav><main>'+''.join(body)+'</main></html>')
for args in [('USER_MANUAL.md','Nami_Grant_User_Manual','manual.html'),('QUICKSTART.md','Nami_Grant_Quick_Start','quickstart.html')]:generate(*args)
