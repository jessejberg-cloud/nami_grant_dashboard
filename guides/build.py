# Builds the two suite-wide guides in guides/dist/ from each program's own QUICKSTART.md and
# USER_MANUAL.md, so the per-program files stay the single source:
#   NAMI_Dashboard_Suite_Quick_Start_Guides.{md,html,docx}
#   NAMI_Dashboard_Suite_User_Manuals.{md,html,docx}
# Then `node guides/print_pdfs.mjs` prints the PDFs from the html. Run both after changing any guide.
from pathlib import Path
from html import escape
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
HERE=Path(__file__).resolve().parent
DIST=HERE/'dist'
DATE='2026-10-05'
PROGRAMS=[('Grant Radar',ROOT/'grant-radar'),('Grant Dashboard',ROOT),('Volunteers & Events',ROOT/'volunteers-and-events')]
NAMES={n for n,_ in PROGRAMS}

def demote(md,title):
    """One program's guide as a section: its H1 becomes an H2 with the program's name, deeper headings move down one."""
    lines=md.splitlines()
    out=[]; skip_contents=False
    for i,line in enumerate(lines):
        if i==0 and line.startswith('# '): out.append('## '+title); continue
        if re.match(r'^## Contents\s*$',line): skip_contents=True; continue
        if skip_contents:
            if line.startswith('## '): skip_contents=False
            else: continue
        if line.startswith('#'): out.append('#'+line); continue
        out.append(line)
    return '\n'.join(out).strip()+'\n'

def inline_html(s):
    s=escape(s)
    s=re.sub(r'`([^`]+)`',r'<code>\1</code>',s)
    s=re.sub(r'\*\*([^*]+)\*\*',r'<strong>\1</strong>',s)
    s=re.sub(r'(?<![*\w])\*([^*\n]+)\*(?![*\w])',r'<em>\1</em>',s)
    s=re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)',r'<a href="\2">\1</a>',s)
    s=re.sub(r'^_(.+)_$',r'<em>\1</em>',s)
    return s

def md_html(text,title):
    out=[]; in_ul=False; in_ol=False; in_code=False; toc=[]
    for raw in text.splitlines():
        line=raw.rstrip()
        if line.startswith('```'):
            if in_code: out.append('</code></pre>'); in_code=False
            else: out.append('<pre><code>'); in_code=True
            continue
        if in_code: out.append(escape(line)+'\n'); continue
        if line.startswith('- '):
            if not in_ul: out.append('<ul>'); in_ul=True
            out.append('<li>'+inline_html(line[2:])+'</li>'); continue
        if re.match(r'^\d+\. ',line):
            if not in_ol: out.append('<ol start="'+re.match(r'^(\d+)\. ',line).group(1)+'">'); in_ol=True
            out.append('<li>'+inline_html(re.sub(r'^\d+\. ','',line))+'</li>'); continue
        if in_ul: out.append('</ul>'); in_ul=False
        if in_ol: out.append('</ol>'); in_ol=False
        if not line: continue
        m=re.match(r'^(#{1,4})\s+(.*)$',line)
        if m:
            level=len(m.group(1)); ident=re.sub(r'[^a-z0-9]+','-',m.group(2).lower()).strip('-')
            if level==2: toc.append((ident,m.group(2)))
            cls=' class="program"' if level==2 and m.group(2) in NAMES else ''
            out.append(f'<h{level} id="{ident}"{cls}>{inline_html(m.group(2))}</h{level}>')
        else: out.append('<p>'+inline_html(line)+'</p>')
    if in_ul: out.append('</ul>')
    if in_ol: out.append('</ol>')
    body='\n'.join(out)
    nav='<nav class="toc"><strong>In this guide</strong><ol>'+''.join(f'<li><a href="#{i}">{escape(t)}</a></li>' for i,t in toc)+'</ol></nav>'
    body=re.sub(r'(</p>)',r'\1'+nav,body,count=1)
    return f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{escape(title)}</title><style>body{{max-width:860px;margin:0 auto;padding:36px 22px 80px;font:17px/1.6 Arial;color:#1e2a28;background:#fffefa}}h1{{font-size:2.1rem;line-height:1.15}}h2{{margin-top:3rem;border-top:3px solid #2d6a64;padding-top:1rem;font-size:1.7rem}}h2.program{{page-break-before:always}}h3{{margin-top:1.8rem;font-size:1.25rem}}h4{{margin-top:1.3rem;font-size:1.05rem}}a{{color:#2d6a64}}code,pre{{background:#e9ede7;padding:.15rem .3rem}}pre{{overflow:auto;padding:1rem}}li{{margin:.3rem 0}}.toc{{background:#e9ede7;border-radius:12px;padding:14px 18px;margin:1.5rem 0}}.toc ol{{margin:.4rem 0 0}}@media(max-width:600px){{body{{padding:22px 16px 60px}}h1{{font-size:1.8rem}}}}@media print{{h2{{border-top:0}}.toc a{{color:inherit;text-decoration:none}}}}</style></head><body>{body}</body></html>'''

def configure(doc,title):
    sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
    st=doc.styles
    st['Normal'].font.name='Arial'; st['Normal'].font.size=Pt(10.5); st['Normal'].font.color.rgb=RGBColor(0,0,0); st['Normal'].paragraph_format.space_after=Pt(6); st['Normal'].paragraph_format.line_spacing=1.08
    for name,size in [('Title',27),('Heading 1',18),('Heading 2',14),('Heading 3',11.5)]:
        s=st[name]; s.font.name='Arial'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0); s.paragraph_format.keep_with_next=True; s.paragraph_format.space_before=Pt(11); s.paragraph_format.space_after=Pt(5)
    from docx.oxml.ns import qn
    ppr=st['Title']._element.get_or_add_pPr(); b=ppr.find(qn('w:pBdr'))
    if b is not None: ppr.remove(b)
    doc.add_paragraph(style='Title').add_run(title)

def add_inline(par,text):
    text=re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)',r'\1',text)
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|(?<![*\w])\*[^*\n]+\*(?![*\w]))',text):
        if part.startswith('**') and part.endswith('**'): par.add_run(part[2:-2]).bold=True
        elif part.startswith('`') and part.endswith('`'): r=par.add_run(part[1:-1]); r.font.name='Courier New'; r.font.size=Pt(9)
        elif part.startswith('*') and part.endswith('*') and len(part)>2: par.add_run(part[1:-1]).italic=True
        else: par.add_run(part)

def md_docx(text,out,title,footer_text):
    from docx.enum.text import WD_BREAK
    doc=Document(); configure(doc,title)
    lines=text.splitlines()[1:]; i=0
    while i<len(lines):
        line=lines[i].rstrip()
        if not line: i+=1; continue
        if line.startswith('```'):
            code=[]; i+=1
            while i<len(lines) and not lines[i].startswith('```'): code.append(lines[i]); i+=1
            p=doc.add_paragraph(); r=p.add_run('\n'.join(code)); r.font.name='Courier New'; r.font.size=Pt(8); i+=1; continue
        m=re.match(r'^(#{2,4})\s+(.*)$',line)
        if m:
            level=len(m.group(1))-1
            if level==1 and m.group(2) in NAMES: doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_paragraph(m.group(2),style=f'Heading {level}'); i+=1; continue
        if line.startswith('- '):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.first_line_indent=Inches(-.2); add_inline(p,'•  '+line[2:]); p.add_run().add_break(); i+=1; continue
        mo=re.match(r'^(\d+)\. ',line)
        if mo:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.32); p.paragraph_format.first_line_indent=Inches(-.28); add_inline(p,f'{mo.group(1)}.  '+line[len(mo.group(0)):]); p.add_run().add_break(); i+=1; continue
        p=doc.add_paragraph(); m2=re.match(r'^_(.+)_$',line)
        if m2: p.add_run(m2.group(1)).italic=True
        else: add_inline(p,line)
        i+=1
    f=doc.sections[0].footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; f.add_run(footer_text)
    for r in f.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor(90,105,110)
    doc.core_properties.title=title; doc.core_properties.author='NAMI Dashboard Suite'; doc.save(out)

def build(kind,title,lead,filename):
    overview=(HERE/'overview.md').read_text()
    parts=['# '+title,'',f'_NAMI Dashboard Suite, {DATE}._','',lead,'',overview]
    for name,folder in PROGRAMS:
        parts.append(demote((folder/kind).read_text(),name))
    md='\n'.join(parts)
    (DIST/(filename+'.md')).write_text(md)
    (DIST/(filename+'.html')).write_text(md_html(md,title))
    md_docx(md,DIST/(filename+'.docx'),title,f'NAMI Dashboard Suite · {DATE}')
    print('wrote',filename,'(.md .html .docx)')

def main():
    DIST.mkdir(exist_ok=True)
    build('QUICKSTART.md','NAMI Dashboard Suite: Quick Start Guides',
          'Three short guides in one place, one for each program. Read the first page, then the guide for the program you were sent. Each program also has these same words inside it, under Help.',
          'NAMI_Dashboard_Suite_Quick_Start_Guides')
    build('USER_MANUAL.md','NAMI Dashboard Suite: User Manuals',
          'The three full manuals in one place, one part for each program. Start with the page on how they fit together, then go to the part for the program you use. Each program also has its own manual inside it, under Help.',
          'NAMI_Dashboard_Suite_User_Manuals')

if __name__=='__main__': main()
