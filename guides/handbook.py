# Builds the handbook: one document, for the binder and for sharing, in guides/dist/:
#   NAMI_Dashboard_Suite_Handbook.html  (printed to PDF, with page numbers, by guides/print_pdfs.mjs)
#   NAMI_Dashboard_Suite_Handbook.docx  (Word, to edit)
# Chapters, leadership first: 1 Overview and the plan, 2 Implementation, 3 Quick start guides,
# 4 User manuals, Appendix: the emails. Every chapter is read from its single source, so nothing
# is copied by hand: overview.md, the one-page plan (docs/project-map.html, as a picture made by
# guides/map_png.mjs), implementation.md, each program's QUICKSTART.md and USER_MANUAL.md, EMAILS.md.
# Order: node guides/map_png.mjs && python3 guides/handbook.py && node guides/print_pdfs.mjs
from pathlib import Path
from html import escape
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from build import ROOT, HERE, DIST, DATE, PROGRAMS, demote, inline_html, add_inline

TITLE='NAMI Dashboard Suite'
DATE_LONG='October 5, 2026'
MAP_PNG=DIST/'project-map.png'

def body_of(md):
    """A source file without its own title line: the tab page carries the title."""
    lines=md.splitlines()
    if lines and lines[0].startswith('# '): lines=lines[1:]
    return '\n'.join(lines).strip()+'\n'

def chapters():
    plan='## The plan on one page\n\n[[MAP]]\n'
    return [
      ('Chapter 1','Overview and the plan','What the three programs are, how they fit together, and the two phases on one page. For everyone; read first.',
       (HERE/'overview.md').read_text()+'\n'+plan),
      ('Chapter 2','Implementation','How the rollout runs week by week, who does what, the weekly routine, privacy, backups, the risks and what we do about each, and phase two. For the Executive Director, the owners and Jesse.',
       body_of((HERE/'implementation.md').read_text())),
      ('Chapter 3','Quick start guides','A few pages for each program: open it, the first things to do, and what to know. For the person using that program.',
       '\n'.join(demote((f/'QUICKSTART.md').read_text(),n) for n,f in PROGRAMS)),
      ('Chapter 4','User manuals','The full manual for each program. The same words are inside each program under Help.',
       '\n'.join(demote((f/'USER_MANUAL.md').read_text(),n) for n,f in PROGRAMS)),
      ('Appendix','The emails','The overview email to the Executive Director, and the three program emails to forward to each owner.',
       body_of((HERE/'EMAILS.md').read_text())),
    ]

# ---------------- HTML ----------------
def table_html(rows):
    head,*rest=rows
    return '<table><thead><tr>'+''.join(f'<th>{inline_html(c)}</th>' for c in head)+'</tr></thead><tbody>'+''.join('<tr>'+''.join(f'<td>{inline_html(c)}</td>' for c in r)+'</tr>' for r in rest)+'</tbody></table>'

def cells(line): return [c.strip() for c in line.strip().strip('|').split('|')]

def md_to_html(text,appendix=False):
    out=[]; lines=text.splitlines(); i=0; lst=None
    def close():
        nonlocal lst
        if lst: out.append(f'</{lst}>'); lst=None
    while i<len(lines):
        line=lines[i].rstrip()
        if line.startswith('|') and i+1<len(lines) and re.match(r'^\|[\s:|-]+\|$',lines[i+1].strip()):
            close(); rows=[cells(line)]; i+=2
            while i<len(lines) and lines[i].startswith('|'): rows.append(cells(lines[i])); i+=1
            out.append(table_html(rows)); continue
        if line=='[[MAP]]':
            close(); out.append('<figure class="map"><img src="project-map.png" alt="The plan on one page: phase one and phase two, the steps under each, who does what, and what comes later."></figure>'); i+=1; continue
        if line.strip()=='---': close(); out.append('<hr>'); i+=1; continue
        if line.startswith('- '):
            if lst!='ul': close(); out.append('<ul>'); lst='ul'
            out.append('<li>'+inline_html(line[2:])+'</li>'); i+=1; continue
        mo=re.match(r'^(\d+)\. ',line)
        if mo:
            if lst!='ol': close(); out.append(f'<ol start="{mo.group(1)}">'); lst='ol'
            out.append('<li>'+inline_html(line[len(mo.group(0)):])+'</li>'); i+=1; continue
        close()
        if not line: i+=1; continue
        m=re.match(r'^(#{2,4})\s+(.*)$',line)
        if m:
            lvl=len(m.group(1)); t=m.group(2)
            cls=' class="newpage"' if lvl==2 and (t in {n for n,_ in PROGRAMS} or (appendix and re.match(r'^\d\.',t))) else ''
            out.append(f'<h{lvl}{cls}>{inline_html(t)}</h{lvl}>'); i+=1; continue
        out.append('<p>'+inline_html(line)+'</p>'); i+=1
    close(); return '\n'.join(out)

CSS='''
@page{size:Letter}
body{margin:0;font:11pt/1.5 Arial,Helvetica,sans-serif;color:#1e2a28;background:#fff}
.page{max-width:7in;margin:0 auto;padding:0 0 .2in}
h2{font-size:17pt;margin:22pt 0 6pt;color:#1e2a28;border-bottom:2px solid #2d6a64;padding-bottom:3pt;break-after:avoid}
h3{font-size:13pt;margin:16pt 0 4pt;break-after:avoid}
h4{font-size:11.5pt;margin:12pt 0 3pt;break-after:avoid}
p{margin:0 0 7pt}ul,ol{margin:0 0 8pt;padding-left:1.4em}li{margin:2pt 0}
a{color:#2d6a64;text-decoration:none}code{background:#e9ede7;padding:0 3px;font-size:10pt}
table{border-collapse:collapse;width:100%;margin:6pt 0 12pt;font-size:9.6pt;break-inside:auto}
th,td{border:1px solid #c9d2c8;padding:5pt 6pt;text-align:left;vertical-align:top}
th{background:#2d6a64;color:#fff;font-weight:700}
tr{break-inside:avoid}
hr{border:0;border-top:1px dashed #c9d2c8;margin:14pt 0}
.newpage{break-before:page}
.cover{height:9.4in;display:flex;flex-direction:column;justify-content:center;break-after:page}
.cover .org{letter-spacing:.12em;text-transform:uppercase;color:#56645f;font-weight:700;font-size:10pt}
.cover h1{font-size:40pt;line-height:1.05;margin:8pt 0 4pt;color:#2d6a64}
.cover .sub{font-size:20pt;margin:0 0 30pt}
.cover .meta{color:#56645f;font-size:11pt;line-height:1.7}
.contents{break-after:page}.contents h2{border:0;font-size:22pt}
.contents ol{list-style:none;padding:0}.contents li{margin:0 0 14pt;padding:10pt 12pt;border-left:4px solid #2d6a64;background:#f2f4f0}
.contents li b{display:block;font-size:12.5pt}.contents li span{color:#56645f;font-size:10pt}
.tab{height:9.4in;display:flex;flex-direction:column;justify-content:center;break-before:page;break-after:page;border-left:10px solid #2d6a64;padding-left:.4in}
.tab .n{letter-spacing:.14em;text-transform:uppercase;color:#56645f;font-weight:700}
.tab h1{font-size:34pt;margin:6pt 0 10pt;color:#1e2a28}
.tab p{font-size:12pt;color:#56645f;max-width:5.2in}
.map{margin:0;break-before:page;break-inside:avoid}.map img{width:100%;border:1px solid #d6ddd5}
@media screen{body{background:#e9ede7}.page{background:#fff;padding:.6in .7in;margin:20px auto}}
'''

def build_html():
    ch=chapters()
    h=[f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{TITLE} handbook</title><style>{CSS}</style></head><body><div class="page">']
    h.append(f'<section class="cover"><div class="org">NAMI Southeast Wisconsin</div><h1>{TITLE}</h1><div class="sub">Handbook</div><div class="meta">The plan, how we roll it out, and the guides for all three programs<br>Phase one: testing · {DATE_LONG}<br>Prepared by Jesse Jonesberg, MSW, LCSW, CPS<br><a href="https://www.intrinsicchange.com">www.intrinsicchange.com</a></div></section>')
    h.append('<section class="contents"><h2>What is in this handbook</h2><ol>'+''.join(f'<li><b>{escape(a)} · {escape(b)}</b><span>{escape(c)}</span></li>' for a,b,c,_ in ch)+'</ol><p>Behind each tab is one chapter. The same handbook, with the same words, is on the NAMI landing page, where it is kept current.</p></section>')
    for a,b,c,md in ch:
        h.append(f'<section class="tab"><div class="n">{escape(a)}</div><h1>{escape(b)}</h1><p>{escape(c)}</p></section>')
        h.append(md_to_html(md,appendix=(a=='Appendix')))
    h.append('</div></body></html>')
    (DIST/'NAMI_Dashboard_Suite_Handbook.html').write_text('\n'.join(h))

# ---------------- Word ----------------
def page_number_footer(sec,text):
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text+' · page '); r.font.size=Pt(8); r.font.color.rgb=RGBColor(90,105,110)
    r2=p.add_run(); r2.font.size=Pt(8)
    for tag,val in [('begin',None),(None,'PAGE'),('end',None)]:
        if tag: e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),tag)
        else: e=OxmlElement('w:instrText'); e.set(qn('xml:space'),'preserve'); e.text=val
        r2._r.append(e)

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); s=OxmlElement('w:shd'); s.set(qn('w:val'),'clear'); s.set(qn('w:fill'),fill); tcPr.append(s)

def page_break(doc): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def md_to_docx(doc,text,appendix=False):
    lines=text.splitlines(); i=0
    names={n for n,_ in PROGRAMS}
    while i<len(lines):
        line=lines[i].rstrip()
        if line.startswith('|') and i+1<len(lines) and re.match(r'^\|[\s:|-]+\|$',lines[i+1].strip()):
            rows=[cells(line)]; i+=2
            while i<len(lines) and lines[i].startswith('|'): rows.append(cells(lines[i])); i+=1
            t=doc.add_table(rows=len(rows),cols=len(rows[0])); t.style='Table Grid'
            for ri,row in enumerate(rows):
                for ci,val in enumerate(row):
                    c=t.cell(ri,ci); p=c.paragraphs[0]; add_inline(p,val)
                    for r in p.runs:
                        r.font.size=Pt(9)
                        if ri==0: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
                    if ri==0: shade(c,'2D6A64')
            doc.add_paragraph(); continue
        if line=='[[MAP]]':
            if MAP_PNG.exists(): page_break(doc); doc.add_picture(str(MAP_PNG),width=Inches(6.9))
            i+=1; continue
        if not line or line.strip()=='---': i+=1; continue
        m=re.match(r'^(#{2,4})\s+(.*)$',line)
        if m:
            lvl=len(m.group(1))-1; t=m.group(2)
            if lvl==1 and (t in names or (appendix and re.match(r'^\d\.',t))): page_break(doc)
            doc.add_paragraph(t,style=f'Heading {lvl}'); i+=1; continue
        if line.startswith('- '):
            p=doc.add_paragraph(style='List Bullet'); add_inline(p,line[2:]); i+=1; continue
        mo=re.match(r'^(\d+)\. ',line)
        if mo:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.32); p.paragraph_format.first_line_indent=Inches(-.26); add_inline(p,f'{mo.group(1)}.  '+line[len(mo.group(0)):]); i+=1; continue
        p=doc.add_paragraph(); m2=re.match(r'^_(.+)_$',line)
        if m2: p.add_run(m2.group(1)).italic=True
        else: add_inline(p,line)
        i+=1

def build_docx():
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    for side in ('top_margin','bottom_margin'): setattr(sec,side,Inches(.75))
    for side in ('left_margin','right_margin'): setattr(sec,side,Inches(.8))
    st=doc.styles
    st['Normal'].font.name='Arial'; st['Normal'].font.size=Pt(10.5); st['Normal'].paragraph_format.space_after=Pt(6)
    for name,size in [('Title',36),('Heading 1',17),('Heading 2',13),('Heading 3',11.5)]:
        s=st[name]; s.font.name='Arial'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0x1E,0x2A,0x28); s.paragraph_format.keep_with_next=True; s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(5)
    ppr=st['Title']._element.get_or_add_pPr(); b=ppr.find(qn('w:pBdr'))
    if b is not None: ppr.remove(b)
    st['Title'].font.color.rgb=RGBColor(0x2D,0x6A,0x64)
    ch=chapters()
    # cover
    for _ in range(8): doc.add_paragraph()
    p=doc.add_paragraph(); r=p.add_run('NAMI SOUTHEAST WISCONSIN'); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor(0x56,0x64,0x5F)
    doc.add_paragraph(TITLE,style='Title')
    p=doc.add_paragraph(); r=p.add_run('Handbook'); r.font.size=Pt(20)
    doc.add_paragraph()
    for t in ['The plan, how we roll it out, and the guides for all three programs',f'Phase one: testing · {DATE_LONG}','Prepared by Jesse Jonesberg, MSW, LCSW, CPS','www.intrinsicchange.com']:
        p=doc.add_paragraph(); r=p.add_run(t); r.font.color.rgb=RGBColor(0x56,0x64,0x5F)
    page_break(doc)
    # contents
    doc.add_paragraph('What is in this handbook',style='Heading 1')
    for a,b,c,_ in ch:
        p=doc.add_paragraph(); r=p.add_run(f'{a} · {b}'); r.bold=True; r.font.size=Pt(12)
        p=doc.add_paragraph(); r=p.add_run(c); r.font.color.rgb=RGBColor(0x56,0x64,0x5F)
    doc.add_paragraph('Behind each tab is one chapter. The same handbook, with the same words, is on the NAMI landing page, where it is kept current.')
    for a,b,c,md in ch:
        page_break(doc)
        for _ in range(9): doc.add_paragraph()
        p=doc.add_paragraph(); r=p.add_run(a.upper()); r.bold=True; r.font.color.rgb=RGBColor(0x56,0x64,0x5F)
        p=doc.add_paragraph(); r=p.add_run(b); r.bold=True; r.font.size=Pt(32)
        p=doc.add_paragraph(); r=p.add_run(c); r.font.size=Pt(12); r.font.color.rgb=RGBColor(0x56,0x64,0x5F)
        page_break(doc)
        md_to_docx(doc,md,appendix=(a=='Appendix'))
    page_number_footer(sec,f'{TITLE} handbook · {DATE}')
    doc.core_properties.title=f'{TITLE} handbook'; doc.core_properties.author='NAMI Dashboard Suite'
    doc.save(DIST/'NAMI_Dashboard_Suite_Handbook.docx')

if __name__=='__main__':
    DIST.mkdir(exist_ok=True); build_html(); build_docx(); print('wrote NAMI_Dashboard_Suite_Handbook (.html .docx)')
