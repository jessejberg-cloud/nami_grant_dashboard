from pathlib import Path
import re, html
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]

def build(source,stem,web,quick=False):
    text=(ROOT/source).read_text().replace('October 4, 2026','October 5, 2026')
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.65 if quick else .7);sec.left_margin=sec.right_margin=Inches(.75)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10 if quick else 10.5);normal.paragraph_format.space_after=Pt(5 if quick else 6);normal.paragraph_format.line_spacing=1.05
    for n in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        doc.styles[n].font.name='Calibri';doc.styles[n].font.color.rgb=RGBColor(0,0,0)
    doc.styles['Title'].font.size=Pt(22 if quick else 25)
    doc.styles['Heading 1'].font.size=Pt(16);doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
    doc.styles['Heading 2'].font.size=Pt(13)
    for st in doc.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    parts=[]
    for line in text.splitlines():
        if not line.strip():continue
        if line=='<!-- PAGE -->':
            doc.add_page_break();continue
        if line.startswith('# '):
            title=re.sub(r'[^\w\s]','',line[2:]);doc.add_paragraph(title,'Title');parts.append('<h1>'+html.escape(title)+'</h1>')
        elif line.startswith('## '):
            title=re.sub(r'[^\w\s]','',line[3:]);
            if not quick and title=='Starting and understanding storage':doc.add_page_break()
            doc.add_paragraph(title,'Heading 1');parts.append('<h2>'+html.escape(title)+'</h2>')
        elif line.startswith('- '):
            doc.add_paragraph(line[2:],'List Bullet');parts.append('<p>• '+html.escape(line[2:])+'</p>')
        else:
            doc.add_paragraph(line);parts.append('<p>'+html.escape(line)+'</p>')
    footer=sec.footer.paragraphs[0];footer.text='Events and Volunteers 1.0.0    |    Fictional evaluation    |    '
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld);footer.runs[0].font.size=Pt(9)
    doc.save(ROOT/'public'/f'{stem}.docx')
    (ROOT/'public'/web).write_text('<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+html.escape(stem.replace('_',' '))+'</title><style>body{max-width:850px;margin:35px auto;padding:0 24px;color:#172c35;font:17px/1.6 system-ui}h1,h2{line-height:1.2}a{color:#086975}@media print{body{font-size:11pt}}</style><a href="/">Return to dashboard</a>'+''.join(parts)+'</html>')

build('USER_MANUAL.md','Nami_Events_User_Manual','manual.html')
build('QUICKSTART.md','Nami_Events_Quick_Start','quickstart.html',True)
text='\n\n'.join((ROOT/f).read_text() for f in ['ADMIN_HANDOFF.md','INTEGRATION.md','PRESENTER.md'])
parts=[]
for l in text.splitlines():
    if l.startswith('# '):parts.append('<h1>'+html.escape(l[2:])+'</h1>')
    elif l.startswith('## '):parts.append('<h2>'+html.escape(l[3:])+'</h2>')
    elif l.strip():parts.append('<p>'+html.escape(l)+'</p>')
(ROOT/'public/admin.html').write_text('<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Administrator handoff</title><style>body{max-width:850px;margin:35px auto;padding:0 24px;font:17px/1.6 system-ui}</style><a href="/">Return to dashboard</a>'+''.join(parts)+'</html>')
